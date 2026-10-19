# -*- coding: utf-8 -*-
"""
套版结果核验 —— ① 行距分布 ② keepNext 计数 ③ 大纲级别分布 ④ 表格是否超版心
⑤ （可选）落单标题全册扫描：需先导出 PDF，再传入 --pdf。

用法：
    python verify_format.py <docx路径> [--pdf <pdf路径>]

注意：
  * keepNext 判定必须看"元素是否存在"——`<w:keepNext/>` 不带 val 即默认生效，
    若用 el.get(qn('w:val')) 读值会得到 None，从而误判为"未设置"。
  * 落单标题扫描有误报：目录条目（含省略号+页码行）、表格内数字行会命中标题正则，
    需人工排除。
"""
import os, re, sys, io
from collections import Counter

from docx import Document
from docx.oxml.ns import qn

RE_HEAD = re.compile(
    r'^(第[一二三四五六七八九十]{1,3}章[\s\u3000]|'
    r'\d{1,2}\.\d{1,2}\.\d{1,2}\s|'
    r'\d{1,2}\.\d{1,2}[\s\u3000]|'
    r'（[一二三四五六七八九十]{1,3}）|'
    r'表\s?\d+\s?[—\-–]\s?\d+)')


def inspect_docx(path, log=print):
    doc = Document(path)
    ls, ol, miss = Counter(), Counter(), 0
    kn = kn_nonempty = 0
    for p in doc.paragraphs:
        ppr = p._element.find(qn('w:pPr'))
        if ppr is None:
            continue
        sp = ppr.find(qn('w:spacing'))
        if sp is not None:
            ls[(sp.get(qn('w:line')), sp.get(qn('w:lineRule')))] += 1
        o = ppr.find(qn('w:outlineLvl'))
        ol[o.get(qn('w:val')) if o is not None else None] += 1
        if ppr.find(qn('w:keepNext')) is not None:      # 只看元素是否存在
            kn += 1
            if p.text.strip():
                kn_nonempty += 1
    # 表格超版心
    sec = doc.sections[0]
    avail = (sec.page_width.cm - sec.left_margin.cm - sec.right_margin.cm)
    over = 0
    for tb in doc.tables:
        try:
            w = sum((c.width.cm if c.width is not None else 0) for c in tb.columns)
        except Exception:
            w = None
        if w and w > avail + 0.01:
            over += 1

    log("  page  %.2f×%.2f cm   margins L/R %.2f/%.2f  T/B %.2f/%.2f cm"
        % (sec.page_width.cm, sec.page_height.cm,
           sec.left_margin.cm, sec.right_margin.cm,
           sec.top_margin.cm, sec.bottom_margin.cm))
    log("  行距分布 (line,lineRule) -> %s     [('360','auto') = 1.5倍]" % dict(ls))
    log("  大纲级别分布 -> %s" % dict(ol))
    log("  keepNext 段落 = %d（非空 %d）" % (kn, kn_nonempty))
    log("  表格 %d 张；列宽合计超版心(%.2fcm) = %d 张" % (len(doc.tables), avail, over))
    return dict(line_spacing=dict(ls), outline=dict(ol), keep_next=kn,
                tables=len(doc.tables), tables_over=over, avail=avail)


def scan_orphan_headings(pdf, log=print):
    """逐页看最后一个非空行是否为标题 → 疑似落单。"""
    import pymupdf
    d = pymupdf.open(pdf)
    hits = []
    for pno, page in enumerate(d, 1):
        lines = [ln.strip() for ln in page.get_text().splitlines() if ln.strip()]
        if not lines:
            continue
        last = lines[-1]
        # 排除目录条目（省略号 + 页码）
        if re.search(r'[.·]{4,}\s*\d+$', last):
            continue
        if len(last) <= 40 and RE_HEAD.match(last):
            hits.append((pno, last))
    d.close()
    log("  页末疑似落单标题 = %d 处（须人工排除表格内数字行）" % len(hits))
    for pno, t in hits:
        log("     p%-4d %s" % (pno, t))
    return hits


if __name__ == '__main__':
    argv = sys.argv[1:]
    pdf = None
    if '--pdf' in argv:
        i = argv.index('--pdf')
        pdf = argv[i + 1] if i + 1 < len(argv) else None
        del argv[i:i + 2]
    paths = [a for a in argv if not a.startswith('--')]
    if not paths:
        print(__doc__); sys.exit(1)
    for p in paths:
        print("=" * 72); print(os.path.basename(p))
        inspect_docx(p)
    if pdf:
        print("=" * 72); print("落单标题扫描：%s" % os.path.basename(pdf))
        scan_orphan_headings(pdf)
