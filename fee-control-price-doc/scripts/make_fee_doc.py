# -*- coding: utf-8 -*-
r"""
二类费用控制价 · 三件套生成（封皮 / 编制说明 / 计费清单）
用法：改 CONFIG 后直接运行。母本路径与输出路径见 CONFIG。
依赖：python-docx、openpyxl（托管 venv 已装）
       C:\Users\lihaoyuan\.workbuddy\binaries\python\envs\default\Scripts\python.exe
"""
import os, shutil, copy
from docx import Document
from docx.text.paragraph import Paragraph
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

# =================== CONFIG ===================
TPL = r"<项目工作区>\医废项目\二类费用控制价\可研费"
OUT = r"<项目工作区>\XXXX\二类费用控制价\XXXX费"

PROJ = "XX项目名称"                       # 项目名称
OWNER = "XX建设单位"                      # 建设单位
SITE = "XX省XX市XX县"                     # 工程地点
SCOPE = PROJ + "编制费"                   # 编制范围
FEE_UPPER = "陆万伍仟元整"                # 大写
FEE_LOWER = "65000.00"                    # 小写
LIST_TITLE = PROJ + "编制费"              # 计费清单表题

# 编制说明 · 三、编制依据及相关费用计取（母本结构为 4 条，建议 ≤5 条以保证单页）
YIJU = [
    "1、" + PROJ + "相关文件；",
    "2、《国家发展改革委关于进一步放开建设项目专业服务价格的通知》（发改价格〔2015〕299号）：前期工作咨询费实行市场调节价，收费标准由委托双方协商确定；",
    "3、《中国工程咨询协会关于工程咨询服务（境内）人工成本要素信息调查情况的通报》（中咨协政〔2015〕46号）：参照“规划咨询”口径按人工成本（人日）法计取，综合计算系数取2.44；",
    "4、《山西省发展和改革委员会投资咨询评估管理办法》附件1：1亿元以下“项目实施方案”评估费标准为4万元/件，本次编制费高于同级评估费，工作量及复杂程度较高；",
    "5、其他计算依据参考现行的文件。",
]

# 计费清单数据：(序号, 工作阶段, 人员类别, 人日单价, 人日数, 金额, 备注)
# 同阶段的第 2 行序号/阶段填 None，脚本自动纵向合并
ROWS = [
    (1, "项目调研与基础数据采集", "高级职称及以上", 2210, 2, 4420, ""),
    (None, None, "中级职称", 1730, 3, 5190, ""),
    (2, "成本测算与方案编制", "高级职称及以上", 2210, 3, 6630, ""),
    (None, None, "中级职称", 1730, 3, 5190, ""),
    (3, "内部审核与专家评审", "外部高级专家", 3000, 1, 3000, ""),
    (None, None, "高级职称及以上", 2210, 1, 2210, ""),
    (4, "直接人工成本小计", "——", "——", 13, 26640, "人日单价及人员类别依据中咨协政〔2015〕46号“规划咨询”口径"),
    (5, "人工日法综合计算系数", "——", 2.44, "——", "——", "同上"),
    (6, "编制费＝人工成本×综合计算系数", "26640×2.44", "——", "——", 65001.60, ""),
    (7, "取定编制费", "取整至百元", "——", "——", 65000.00, "大写：" + FEE_UPPER),
]
MERGE_PAIRS = [(3, 4), (5, 6), (7, 8)]     # 序号/阶段纵向合并的行对
# ==============================================


def set_para_text(p, text):
    """整段文本写入 run0，清空其余 run —— 保住母本 run 级字体格式"""
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ''


def clone_after(p, text):
    new_el = copy.deepcopy(p._element)
    p._element.addnext(new_el)
    np = Paragraph(new_el, p._parent)
    set_para_text(np, text)
    return np


def build_cover():
    shutil.copy(os.path.join(TPL, "封皮.docx"), os.path.join(OUT, "封皮.docx"))
    d = Document(os.path.join(OUT, "封皮.docx"))
    hit = []
    for p in d.paragraphs:
        t = p.text
        if t.startswith("工程名称："):
            set_para_text(p, "工程名称：" + PROJ); hit.append("name")
        elif t.startswith("内容："):
            set_para_text(p, "内容：" + SCOPE); hit.append("scope")
        elif t.startswith("最高投标限价"):
            for r in p.runs:
                if "元整" in r.text:
                    r.text = FEE_UPPER
            hit.append("upper")
        elif t.startswith("(小写)") or t.startswith("（小写）"):
            for r in p.runs:
                if any(c.isdigit() for c in r.text):
                    r.text = FEE_LOWER
            hit.append("lower")
    assert len(hit) == 4, "封皮定位不全：%s（母本文案被改动？）" % hit
    d.save(os.path.join(OUT, "封皮.docx"))
    print("[OK] 封皮.docx")


def build_shuoming():
    shutil.copy(os.path.join(TPL, "编制说明.docx"), os.path.join(OUT, "编制说明.docx"))
    d = Document(os.path.join(OUT, "编制说明.docx"))

    def find(pred):
        for p in d.paragraphs:
            if pred(p.text):
                return p
        return None

    p_name = find(lambda t: t.strip().startswith("1、工程名称："))
    p_owner = find(lambda t: t.strip().startswith("2、建设单位："))
    p_site = find(lambda t: t.strip().startswith("3、工程地点："))
    p_scope = find(lambda t: "费" in t and not t.strip()[0].isdigit() and "：" not in t)
    p_dep1 = find(lambda t: t.strip().startswith("1、") and ("文件" in t))
    assert all([p_name, p_owner, p_site, p_scope, p_dep1]), "编制说明关键段落定位失败"

    set_para_text(p_name, "1、工程名称：" + PROJ)
    set_para_text(p_owner, "2、建设单位：" + OWNER)
    set_para_text(p_site, "3、工程地点：" + SITE)
    set_para_text(p_scope, SCOPE)
    set_para_text(p_dep1, YIJU[0])

    allp = list(d.paragraphs)
    i0 = next(k for k, x in enumerate(allp) if x._element is p_dep1._element)
    follow = []
    for j in range(i0 + 1, len(allp)):
        if allp[j].text.strip():
            follow.append(allp[j])
        if len(follow) == 3:            # 母本除第 1 条外还有 3 条
            break
    assert len(follow) == 3, "母本依据段数变了"
    set_para_text(follow[0], YIJU[1])
    set_para_text(follow[1], YIJU[2])
    set_para_text(follow[2], YIJU[3])
    last = follow[2]
    for extra in YIJU[4:]:
        last = clone_after(last, extra)

    # 删除尾部空段（防内容增行后多出空白页）
    while True:
        ps = d.paragraphs
        if len(ps) <= 1 or ps[-1].text.strip():
            break
        el = ps[-1]._element
        el.getparent().remove(el)
    assert d.paragraphs and d.paragraphs[-1].text.strip(), "尾部段落异常"

    d.save(os.path.join(OUT, "编制说明.docx"))
    print("[OK] 编制说明.docx")


def build_list():
    wb = Workbook(); ws = wb.active; ws.title = "计费清单"
    thin = Side(style="thin", color="000000")
    BD = Border(left=thin, right=thin, top=thin, bottom=thin)
    CT = Alignment(horizontal="center", vertical="center", wrap_text=True)
    CT_NW = Alignment(horizontal="center", vertical="center")
    LT = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws.merge_cells("A1:G1")
    ws["A1"] = LIST_TITLE
    ws["A1"].font = Font(name="宋体", size=18, bold=True)
    ws["A1"].alignment = CT

    for i, h in enumerate(["序号", "工作阶段", "人员类别", "人日单价\n（元/人日）",
                           "人日数\n（人日）", "金额（元）", "备注"], start=1):
        c = ws.cell(row=2, column=i, value=h)
        c.font = Font(name="宋体", size=11, bold=True); c.alignment = CT; c.border = BD

    r = 3
    for row in ROWS:
        for i, v in enumerate(row, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = Font(name="宋体", size=11)
            c.border = BD
            c.alignment = LT if i == 7 else CT_NW
            if i == 6:
                c.number_format = "#,##0.00"
            elif i == 4 and isinstance(v, float):
                c.number_format = "0.00"
        r += 1

    for a, b in MERGE_PAIRS:                 # 写完数据再合并，否则 MergedCell 只读
        ws.merge_cells(start_row=a, start_column=1, end_row=b, end_column=1)
        ws.merge_cells(start_row=a, start_column=2, end_row=b, end_column=2)

    for k, w in {"A": 6.18, "B": 32, "C": 16, "D": 13, "E": 11, "F": 13, "G": 26}.items():
        ws.column_dimensions[k].width = w
    ws.row_dimensions[1].height = 58
    ws.row_dimensions[2].height = 42
    for rr in range(3, r):
        ws.row_dimensions[rr].height = 36

    ws.page_setup.paperSize = 9                      # A4
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1

    wb.save(os.path.join(OUT, "计费清单.xlsx"))
    print("[OK] 计费清单.xlsx")


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    build_cover()
    build_shuoming()
    build_list()
