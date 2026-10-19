# -*- coding: utf-8 -*-
# [skill: settlement-merge-summary] gap_report.py 自测
"""思路：**反向构造一处已知缺陷**，验证 gap_report.py 能否把它揪出来。

做法（通用，不依赖具体项目）：
 1. 在源表审定列(col8)里找一个「公式行」，其引用集合**全部是叶子行**（即直接引用定额行）
 2. 用外科式改 XML 把该公式重写成"显式列表、去掉最后一个引用项"
    → 这就制造了「父行漏引一行」的缺陷
 3. 同步把项目级合计行(r4)的**缓存值**减去被删行的金额
    （只改缓存就够：gap_report 的 gap 用 r4 缓存算，公式链覆盖用公式算）
 4. 跑 gap_report.py，断言它能报出：
    · 差异 == 被删行金额
    · 漏覆盖行 == {被删行}
    · 缺陷定位 == 被改的那个公式行
 5. 清理临时文件

通过则说明 gap_report.py 的链展开、差异计算、缺陷定位、证据提取均正常。
"""
import os, re, zipfile, subprocess, shutil
import openpyxl
from decimal import Decimal as D

# ============ 配置区 ============
BASE = r"<项目工作区>\审核过程资料"
SRC  = os.path.join(BASE, "工程结算书8-28 - 副本.xlsx")
SK   = r"C:\Users\lihaoyuan\.workbuddy\skills\settlement-merge-summary\scripts"
WS   = r"D:\Workbuddy文件\2026-07-14-15-47-54"
SRC_SHEET = "Sheet2"
TOP_ROW = 4            # 项目级合计行
# ===============================

TOKEN = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?")
def ci(col):
    n = 0
    for ch in col: n = n * 26 + (ord(ch) - 64)
    return n
def refs_rows(f, col):
    """公式中 col 列引用的行号集合"""
    if not isinstance(f, str) or not f.startswith("="): return set()
    out = set()
    for m in TOKEN.finditer(f):
        c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3), m.group(4)
        if c2 and ci(c1) == col and ci(c2) == col:
            out.update(range(min(r1, int(r2)), max(r1, int(r2)) + 1))
        elif not c2 and ci(c1) == col:
            out.add(r1)
    return out

FAKE = os.path.join(WS, "_selftest_src.xlsx")
TMPP = os.path.join(WS, "_selftest_run.py")
REP  = os.path.join(WS, "_selftest_report.xlsx")
LOG  = os.path.join(WS, "_selftest_report_log.txt")

for p in (FAKE, TMPP, REP, LOG):
    if os.path.exists(p): os.remove(p)

# ---------- 1. 找候选公式行 ----------
wv = openpyxl.load_workbook(SRC, data_only=True)
wf = openpyxl.load_workbook(SRC)
sv, sf = wv[SRC_SHEET], wf[SRC_SHEET]
leaf = {r for r in range(TOP_ROW + 1, sv.max_row + 1)
        if isinstance(sv.cell(r, 6).value, (int, float)) and sv.cell(r, 2).value}

cand = None
for r in range(TOP_ROW + 1, sv.max_row + 1):
    S = sorted(refs_rows(sf.cell(r, 8).value, 8))
    if len(S) >= 3 and all(x in leaf for x in S):
        cand = (r, S); break
assert cand, "未找到『引用集合全为叶子行』的公式行，无法自测"
cand_row, S = cand
DROPPED = S[-1]                       # 将被"漏掉"的那一行
AMT = D(repr(sv.cell(DROPPED, 8).value))
print("构造缺陷：源表 r%d 的审定列公式将漏掉 r%d（金额 %s）" % (cand_row, DROPPED, AMT))
print("  原公式：%s" % sf.cell(cand_row, 8).value)

# ---------- 2. 外科式改 XML ----------
zin = zipfile.ZipFile(SRC)
items = [(i, zin.read(i.filename)) for i in zin.infolist()]
zin.close()
SX = "xl/worksheets/sheet1.xml"
xml = next(d for i, d in items if i.filename == SX).decode("utf-8")

# 2a) 目标行公式 → 显式列表去掉最后一项
m = re.search(r'<c r="H%d"[^>]*>.*?</c>' % cand_row, xml, re.S)
assert m, "XML 中找不到 H%d" % cand_row
old = m.group(0)
s_attr = re.search(r's="(\d+)"', old).group(1)
newf = "SUM(" + ",".join("H%d" % x for x in S[:-1]) + ")"
newv = float(sum(D(repr(sv.cell(x, 8).value)) for x in S[:-1]))
xml = xml.replace(old, '<c r="H%d" s="%s"><f>%s</f><v>%s</v></c>'
                  % (cand_row, s_attr, newf, ("%.10f" % newv).rstrip("0").rstrip(".")), 1)
print("  构造后：=SUM(...) 仅 %d 项，去掉 r%d" % (len(S) - 1, DROPPED))

# 2b) 项目级合计行缓存 -= 金额
m4 = re.search(r'<c r="H%d"[^>]*>.*?</c>' % TOP_ROW, xml, re.S)
assert m4, "XML 中找不到 H%d" % TOP_ROW
old4 = m4.group(0)
s4 = re.search(r's="(\d+)"', old4).group(1)
f4 = re.search(r'<f>([^<]*)</f>', old4).group(1)
v4 = D(re.search(r'<v>([^<]*)</v>', old4).group(1))
newv4 = v4 - AMT
xml = xml.replace(old4, '<c r="H%d" s="%s"><f>%s</f><v>%s</v></c>'
                  % (TOP_ROW, s4, f4, ("%.10f" % newv4).rstrip("0").rstrip(".")), 1)
print("  同步 H%d 缓存：%s → %s" % (TOP_ROW, v4, newv4))

with zipfile.ZipFile(FAKE, "w", zipfile.ZIP_DEFLATED) as z:
    for info, data in items:
        z.writestr(info, xml.encode("utf-8") if info.filename == SX else data)

# ---------- 3. 跑 gap_report.py ----------
t = open(os.path.join(SK, "gap_report.py"), encoding="utf-8").read()
OLD_SRC = 'SRC    = os.path.join(BASE, "工程结算书8-28 - 副本.xlsx")'
assert OLD_SRC in t, "gap_report.py 里的 SRC 行变了，自测需同步更新"
t = t.replace(OLD_SRC, 'SRC    = r"%s"' % FAKE)
# 注意：替换串里含 Windows 路径的反斜杠，必须用 lambda 避免 re 解析转义（否则 \W 报 bad escape）
t = re.sub(r'REPORT = r"[^"]+"', lambda m: 'REPORT = r"%s"' % REP, t, count=1)
assert 'REPORT = r"%s"' % REP in t, "REPORT 替换失败"
open(TMPP, "w", encoding="utf-8").write(t)

print()
print("=" * 74)
r = subprocess.run([r"C:\Users\lihaoyuan\AppData\Local\Programs\Python\Python312\python.exe", TMPP],
                   capture_output=True, text=True, encoding="utf-8", errors="replace")
print(r.stdout.strip() or "(无 stdout)")
if r.stderr and "dirname" not in r.stderr:
    print("STDERR:", r.stderr[-500:])

txt = open(LOG, encoding="utf-8").read() if os.path.exists(LOG) else ""

# ---------- 4. 断言 ----------
print()
print("=" * 74)
print("自测断言")
print("=" * 74)
checks = [
    ("差异 == 被删行金额 %.4f" % AMT, ("%.4f" % AMT) in txt and ("差异 ①−②" in txt)),
    ("漏覆盖恰为 r%d" % DROPPED, ("r%d" % DROPPED) in txt and "H 漏、G 有" in txt),
    ("缺陷定位 = r%d" % cand_row, ("分叉点：r%d" % cand_row) in txt),
    ("报出『G 有 H 没有』的行号", ("['r%d']" % DROPPED) in txt),
    ("生成报告文件", os.path.exists(REP)),
]
for name, ok in checks:
    print("  %s %s" % ("✅" if ok else "❌", name))
ok_all = all(o for _, o in checks)

# ---------- 5. 清理 ----------
print()
for p in (FAKE, TMPP):
    if os.path.exists(p): os.remove(p)
print("临时文件已清理（报告与日志保留供查看）：")
print("  %s" % REP)
print("  %s" % LOG)
print()
print("自测通过 ✅" if ok_all else "自测未通过 ❌ —— gap_report.py 需检查")
raise SystemExit(0 if ok_all else 1)
