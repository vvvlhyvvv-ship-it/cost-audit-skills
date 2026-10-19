# -*- coding: utf-8 -*-
# [skill: settlement-merge-summary] 修复 D/E 列工程量精度截断（Decimal 精确聚合、只改工程量列）
# 用法：改顶部 BASE/SRC/DST 路径常量后直接运行；输出写入 .txt 便于 Read 查看。
"""修复 合并汇总 D/E 列（报送工程量 / 审核工程量）的精度截断

问题：build_hebing9_final.py 落盘时用 round(工程量, 3)，把源表聚合量截断到 3 位小数。
      源表存在 0.0125 / 8.34652705061082 这类高精度工程量，截断后
      J 列(=E*F) 与 H 列（源表审定金额原值）产生偏差（最大 40.01 元，Σ 差 2.72 元）。

修法：按同一套四元组主键重新聚合源表，把 D/E 写回完整精度（Decimal 精确求和，保留 10 位）。
      金额列 G/H/I 与合计行不动 —— 它们本就是源表原值，合计不会变。
"""
import openpyxl, re, shutil, os
from collections import defaultdict
from decimal import Decimal

BASE = r"<项目工作区>\审核过程资料"
SRC = os.path.join(BASE, "工程结算书8-28 - 副本.xlsx")
DST = os.path.join(BASE, "工程结算书合并汇总.xlsx")
OUT = r"D:\Workbuddy文件\2026-07-14-15-47-54\hb9_fix_qty_log.txt"
L = []
def P(s=""): L.append(str(s))
D = lambda x: Decimal(repr(x)) if isinstance(x, (int, float)) else Decimal(0)

# ============ 1. 源表聚合（完整精度） ============
def nos(s):
    return "" if s is None else re.sub(r"\s+", "", str(s))

def nu(u):
    if u is None: return ""
    s = re.sub(r"\s+", "", str(u))
    return s.replace("m3", "m³").replace("m2", "m²").replace("hm2", "hm²").replace("km2", "km²")

ALIAS = [("装石碴自卸汽车运输", "石碴运输"), ("挖掘机装石碴", "石碴"),
         ("挖掘机挖装运土", "挖装自卸汽车运土"), ("伸缩缝", "道路伸缩缝")]
SUFFIXES = [
    r"自卸汽车柴油型载重量[\d\.]+t", r"柴油型载重量[\d\.]+t", r"载重量[\d\.]+t",
    r"推土机功率\d+kW", r"功率\d+kW",
    r"1m³挖掘机", r"挖掘机装石碴自卸汽车运输", r"装石碴自卸汽车运输",
    r"运距[0-9.~～\-]+km", r"运距[0-9.~～\-]+", r"运距$",
    r"推土距离.*",
    r"厚度\d+cm", r"增加[\-\d\.]+cm", r"三类土", r"机械夯填", r"其他机械",
    r"φ[\d\.]+mm", r"平段公称直径\d+mm以内", r"（51m）", r"\(51m\)",
    r"机械摊铺压实", r"机械摊铺", r"压实", r"沥青木板",
    r"\[[^\]]*\]", r"（[^）]*[：:][^）]*）", r"\([^)]*[:：][^)]*\)",
]
def nq(s):
    n = nos(s)
    for a, b in ALIAS: n = n.replace(a, b)
    for suf in SUFFIXES: n = re.sub(suf, "", n)
    return re.sub(r"(C25混凝土渠).*", r"\1", n)

ws2 = openpyxl.load_workbook(SRC, data_only=True)["Sheet2"]
agg = defaultdict(lambda: {"sb": Decimal(0), "sh": Decimal(0)})
parent = None
for r in range(4, ws2.max_row + 1):
    nm = ws2.cell(r, 2).value
    if nm is None: continue
    c6 = ws2.cell(r, 6).value
    if isinstance(c6, (int, float)):
        c4, c5 = ws2.cell(r, 4).value, ws2.cell(r, 5).value
        if not isinstance(c4, (int, float)) and not isinstance(c5, (int, float)): continue
        key = (nq(parent) if parent else "", nq(str(nm)), nu(ws2.cell(r, 3).value), round(c6, 2))
        agg[key]["sb"] += D(c4) if isinstance(c4, (int, float)) else Decimal(0)
        agg[key]["sh"] += D(c5) if isinstance(c5, (int, float)) else Decimal(0)
    else:
        parent = str(nm)
P("源表聚合项 %d 个" % len(agg))

# ============ 2. 目标表叶子行 ============
wb_v = openpyxl.load_workbook(DST, data_only=True)
wb_w = openpyxl.load_workbook(DST)
wmv, wmw = wb_v["合并汇总"], wb_w["合并汇总"]

def level_of(seq):
    s = str(seq)
    if s.startswith("（") or s.startswith("("): return 0
    return s.count(".") + 1

stack, items = [], []
for r in range(4, wmv.max_row + 1):
    idx, nm, u = wmv.cell(r, 1).value, wmv.cell(r, 2).value, wmv.cell(r, 3).value
    c4, c5, c6 = wmv.cell(r, 4).value, wmv.cell(r, 5).value, wmv.cell(r, 6).value
    if nm is None: continue
    if str(nm).strip() == "合  计": items.append({"r": r, "kind": "total"}); continue
    if isinstance(c6, (int, float)):
        items.append({"r": r, "nm": str(nm), "u": u, "price": round(c6, 2), "kind": "leaf",
                      "parent": stack[-1][1] if stack else None,
                      "old_c4": c4, "old_c5": c5,
                      "fmt4": wmw.cell(r, 4).number_format, "fmt5": wmw.cell(r, 5).number_format})
    else:
        lv = level_of(idx) if idx is not None else 999
        while stack and stack[-1][0] >= lv: stack.pop()
        stack.append((lv, str(nm)))
        items.append({"r": r, "kind": "header"})

used = set()
def find_match(name_n, unit, price, parent_n):
    k = (parent_n, name_n, unit, price)
    if k in agg: used.add(k); return agg[k]
    if "/" in name_n:
        tot = {"sb": Decimal(0), "sh": Decimal(0)}; hit = False
        for part in name_n.split("/"):
            k2 = (parent_n, nq(part), unit, price)
            if k2 in agg:
                tot["sb"] += agg[k2]["sb"]; tot["sh"] += agg[k2]["sh"]; used.add(k2); hit = True
        if hit: return tot
    return None

leaves = [i for i in items if i.get("kind") == "leaf"]
matched = 0; miss = []
P("目标叶子行 %d 个" % len(leaves))
for it in leaves:
    m = find_match(nq(it["nm"]), nu(it["u"]), it["price"], nq(it["parent"]) if it["parent"] else "")
    if m:
        matched += 1
        it["new_c4"] = float(round(m["sb"], 10))
        it["new_c5"] = float(round(m["sh"], 10))
    else:
        it["new_c4"] = it["old_c4"]; it["new_c5"] = it["old_c5"]; miss.append(it)
P("匹配 %d / %d" % (matched, len(leaves)))
for it in miss: P("  !! 未匹配 r%d %s" % (it["r"], it["nm"]))

# ============ 3. 差异诊断（修复前 vs 修复后） ============
P("")
P("=== D/E 列精度修复明细（仅列差异 > 1e-9 的行）===")
P("  %-4s %-24s %-26s %-26s" % ("行", "定额", "报送量 旧→新", "审核量 旧→新"))
changed = 0
for it in leaves:
    o4, n4, o5, n5 = it["old_c4"], it["new_c4"], it["old_c5"], it["new_c5"]
    d4 = abs(D(o4) - D(n4)) if isinstance(o4, (int, float)) else Decimal(0)
    d5 = abs(D(o5) - D(n5)) if isinstance(o5, (int, float)) else Decimal(0)
    if d4 > Decimal("1e-9") or d5 > Decimal("1e-9"):
        changed += 1
        P("  r%-3d %-24s %-14s→%-14s %-14s→%-14s" % (
            it["r"], it["nm"][:22],
            o4 if o4 is not None else "-", ("%.10g" % n4),
            o5 if o5 is not None else "-", ("%.10g" % n5)))
P("  共 %d 行工程量需要恢复精度" % changed)
P("  典型行的数字格式：D 列='%s'  E 列='%s'" % (leaves[0]["fmt4"], leaves[0]["fmt5"]))

# ============ 4. 落盘 ============
bak = DST.replace(".xlsx", "_备份20260914_修精度前.xlsx")
if not os.path.exists(bak):
    shutil.copy2(DST, bak); P("\n已备份 -> %s" % os.path.basename(bak))

for it in leaves:
    wmw.cell(it["r"], 4, it["new_c4"])
    wmw.cell(it["r"], 5, it["new_c5"])
wb_w.save(DST)
P("[已落盘] %s" % os.path.basename(DST))

# ============ 5. 全表验证 ============
wv2 = openpyxl.load_workbook(DST, data_only=True)["合并汇总"]
P("")
P("=== 修复后验证 ===")
n = 0; s7 = s8 = Decimal(0); bad_j = bad_d = 0
sum_j = sum_e = Decimal(0)
worst = []
for r in range(4, 102):
    e, f, g, h = wv2.cell(r, 5).value, wv2.cell(r, 6).value, wv2.cell(r, 7).value, wv2.cell(r, 8).value
    d = wv2.cell(r, 4).value
    if not (isinstance(f, (int, float)) and isinstance(h, (int, float))): continue
    n += 1; s7 += D(g) if isinstance(g, (int, float)) else Decimal(0); s8 += D(h) if isinstance(h, (int, float)) else Decimal(0)
    if isinstance(e, (int, float)):
        cj = D(e) * D(f); sum_j += cj
        dj = D(h) - cj
        if abs(dj) >= Decimal("0.005"): bad_j += 1; worst.append((r, abs(dj), float(dj)))
    if isinstance(d, (int, float)) and isinstance(g, (int, float)):
        sum_e += D(d) * D(f)
        if abs(D(g) - D(d) * D(f)) >= Decimal("0.005"): bad_d += 1
P("  叶子行 %d" % n)
P("  Σ报送金额 G = %s" % s7)
P("  Σ审核金额 H = %s" % s8)
P("  Σ(E*F)     = %s   （与 ΣH 差 %s）" % (sum_j, s8 - sum_j))
P("  Σ(D*F)     = %s   （与 ΣG 差 %s）" % (sum_e, s7 - sum_e))
P("")
P("  💡 J=E*F 与 H 差异 ≥0.005 的行数：%d  （修复前 10）" % bad_j)
for r, a, d in sorted(worst, key=lambda t: -t[1])[:10]:
    nm = wv2.cell(r, 2).value
    P("     r%-3d %-26s 差 %+.6f" % (r, str(nm)[:24], d))
P("  💡 D*F 与 G 差异 ≥0.005 的行数：%d" % bad_d)
P("")
P("  合计行 r102: 报送=%s  审核=%s  差额=%s" % (wv2.cell(102, 7).value, wv2.cell(102, 8).value, wv2.cell(102, 9).value))
P("  r103 校验值 = %s" % wv2.cell(103, 8).value)
P("  >>> ΣH − r102 = %s    (应 0)" % (s8 - D(wv2.cell(102, 8).value)))
P("  >>> r102 − r103 = %s (应 0 → r104 归零)" % (D(wv2.cell(102, 8).value) - D(wv2.cell(103, 8).value)))

open(OUT, "w", encoding="utf-8").write("\n".join(L))
print("done -> D/E 修复完成, 改动 %d 行" % changed)
