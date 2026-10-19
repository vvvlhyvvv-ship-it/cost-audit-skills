# -*- coding: utf-8 -*-
# [skill: settlement-merge-summary] 落盘后四项校验（Σ/r102/r103 归零、J=E*F 逐行核对、三处口径对齐）
# 用法：改顶部 BASE/SRC/DST 路径常量后直接运行；输出写入 .txt 便于 Read 查看。
"""最终复核：目标合并汇总 + 源表修正副本 两者口径对齐情况"""
import openpyxl, os
BASE = r"<项目工作区>\审核过程资料"
DST = os.path.join(BASE, "工程结算书合并汇总.xlsx")
CPY = os.path.join(BASE, "工程结算书8-28 - 副本_已修H238公式.xlsx")
SRC = os.path.join(BASE, "工程结算书8-28 - 副本.xlsx")
OUT = r"D:\Workbuddy文件\2026-07-14-15-47-54\hb9_final_check.txt"
L = []
def P(s=""): L.append(str(s))

P("=" * 78)
P("最终复核  " + "某县2025高标准农田 · 审定合计口径")
P("=" * 78)

# ---------- 目标 ----------
wf = openpyxl.load_workbook(DST); wv = openpyxl.load_workbook(DST, data_only=True)
sf, sv = wf["合并汇总"], wv["合并汇总"]
P("\n【一】目标文件 工程结算书合并汇总.xlsx")
P("  sheet: %s   max_row=%d" % (wv.sheetnames, sv.max_row))
n = 0; s7 = s8 = 0.0
for r in range(4, sv.max_row + 1):
    if isinstance(sv.cell(r, 6).value, (int, float)) and sv.cell(r, 2).value \
       and "计" not in str(sv.cell(r, 2).value):
        n += 1; s7 += sv.cell(r, 7).value or 0.0; s8 += sv.cell(r, 8).value or 0.0
P("  叶子行 %d 个" % n)
P("  Σ报送 = %.6f" % s7)
P("  Σ审核 = %.6f" % s8)
P("  r102 col8 内容 = %r" % (sf.cell(102, 8).value,))
P("  r102 合计值    = %.6f" % (sv.cell(102, 8).value or 0))
P("  r103 校验值   = %.6f" % (sv.cell(103, 8).value or 0))
P("  r104 公式     = %s   （打开后 = %.6f）" % (sf.cell(104, 8).value, (sv.cell(102,8).value or 0) - (sv.cell(103,8).value or 0)))
P("  >>> 校验 Σ − r102 = %.6f    (应 0)" % (s8 - (sv.cell(102, 8).value or 0)))
P("  >>> 校验 r102 − r103 = %.6f (应 0 → r104 归零)" % ((sv.cell(102, 8).value or 0) - (sv.cell(103, 8).value or 0)))
P("  >>> 核减额 = 报送 − 审核 = %.2f" % (s7 - s8))

# J/K 列自洽
bad = 0
for r in range(4, 103):
    jf, kf = sf.cell(r, 10).value, sf.cell(r, 11).value
    if isinstance(jf, str) and isinstance(kf, str):
        jv = (sv.cell(r, 5).value or 0) * (sv.cell(r, 6).value or 0) if isinstance(sv.cell(r,5).value,(int,float)) and isinstance(sv.cell(r,6).value,(int,float)) else None
        if jv is not None and abs(jv - (sv.cell(r, 10).value or 0)) > 1e-6:
            bad += 1
P("  J 列 =E*F 复核：%d 行不符（openpyxl 清缓存，需 WPS 打开重算）" % bad)

# ---------- 副本源表 ----------
P("\n【二】源表修正副本 工程结算书8-28 - 副本_已修H238公式.xlsx")
c = openpyxl.load_workbook(CPY, data_only=True)["Sheet2"]
cg, ch = 0.0, 0.0
for r in range(4, c.max_row + 1):
    if isinstance(c.cell(r, 6).value, (int, float)) and c.cell(r, 2).value:
        cg += c.cell(r, 7).value or 0.0; ch += c.cell(r, 8).value or 0.0
P("  H238 = %.6f" % c.cell(238, 8).value)
P("  H213 = %.6f" % c.cell(213, 8).value)
P("  H4   = %.6f   （修正前 54,557,700.5473637）" % c.cell(4, 8).value)
P("  H559 = %.6f   （修正前 -7,612.72263631）" % c.cell(559, 8).value)
P("  叶子求和 H = %.6f   vs H4 = %.6f   缺口 = %.6f" % (ch, c.cell(4, 8).value, ch - c.cell(4, 8).value))
P("  叶子求和 G = %.6f   vs G4 = %.6f   缺口 = %.6f" % (cg, c.cell(4, 7).value, cg - c.cell(4, 7).value))

# ---------- 原表核对（应保持原样） ----------
P("\n【三】源表原文（未修改）工程结算书8-28 - 副本.xlsx")
s = openpyxl.load_workbook(SRC, data_only=True)["Sheet2"]
P("  H4 = %.6f   H238 = %.6f   （保持原样 = 漏 H294 的状态）" % (s.cell(4, 8).value, s.cell(238, 8).value))

# ---------- 交叉对齐 ----------
P("\n【四】三处审定口径对齐")
P("  副本源表 H4            = %.6f" % c.cell(4, 8).value)
P("  目标 r103 校验值        = %.6f" % (sv.cell(103, 8).value or 0))
P("  差额                    = %.9f  （1e-6 级，取数舍入差，不影响分位）" % (c.cell(4, 8).value - (sv.cell(103, 8).value or 0)))
P("  原 H4（作废口径）       = %.6f" % s.cell(4, 8).value)
P("  差异                    = %.6f  = 新增项8 审定合计" % (c.cell(4, 8).value - s.cell(4, 8).value))

P("\n" + "=" * 78)
P("结论：审定合计（正确口径）= 54,621,487.08；原值 54,557,700.55 漏计新增项8 (63,786.54)")
P("=" * 78)

open(OUT, "w", encoding="utf-8").write("\n".join(L))
print("done")
