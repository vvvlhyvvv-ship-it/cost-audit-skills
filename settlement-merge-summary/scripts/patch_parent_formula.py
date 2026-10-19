# -*- coding: utf-8 -*-
# [skill: settlement-merge-summary] 外科式修源表父行公式（含「被占用则落副本」双分支、幂等）
# 用法：改顶部 BASE/SRC/DST 路径常量后直接运行；输出写入 .txt 便于 Read 查看。
"""源表 H238 公式修正 —— 副本落盘模式（源表被 WPS/Excel 占用时的替代路径）

背景：源表 Sheet2 的 H238（项2 汇总行）审定列公式漏引 H294（新增项8 马家沟挡土墙
      部分），而报送列 G238 已含 G294 —— 属"插入新行时只改了报送列、漏改审定列"。
      导致 H4（总审定）少算 63,786.5364。

本脚本做两件事，二选一：
  A) 源表可写 → 原地外科式修正（公式 + 同步 4 个缓存值），其余条目/缓存原样保留
  B) 源表被占用 → 输出修正副本「..._已修H238公式.xlsx」，原表不动

外科式改 XML 的理由：openpyxl 存盘会清空全表公式缓存（该表被下游 data_only 读取），
所以只 zipfile 直改 xl/worksheets/sheet1.xml 中的 <f>/<v>。
"""
import zipfile, re, os, shutil

BASE = r"<项目工作区>\审核过程资料"
SRC  = os.path.join(BASE, "工程结算书8-28 - 副本.xlsx")
CPY  = os.path.join(BASE, "工程结算书8-28 - 副本_已修H238公式.xlsx")
OUT  = r"D:\Workbuddy文件\2026-07-14-15-47-54\hb9_patch_src_copy_log.txt"
L = []
def P(s=""): L.append(str(s))

SHEET_XML = "xl/worksheets/sheet1.xml"
OLD_H238_F = "SUM(H239,H247,H252,H261,H269,H277,H286)"
NEW_H238_F = "SUM(H239,H247,H252,H261,H269,H277,H286,H294)"

# ---------- 读源表，动态取 H294（不硬编码） ----------
import openpyxl
wsv = openpyxl.load_workbook(SRC, data_only=True)["Sheet2"]
H294 = wsv.cell(294, 8).value
assert isinstance(H294, (int, float)), "H294 取不到数值"
H238_OLD = wsv.cell(238, 8).value
H213_OLD = wsv.cell(213, 8).value
H4_OLD   = wsv.cell(4, 8).value
H559_OLD = wsv.cell(559, 8).value
H558     = wsv.cell(558, 8).value

# 补齐目标值：一律以"旧缓存 + H294"递推，保证与表内既有口径一致
H238_NEW = H238_OLD + H294
H213_NEW = H213_OLD + H294
H4_NEW   = H4_OLD   + H294
H559_NEW = H4_NEW - H558

P("=== 修正目标（源表 Sheet2）===")
P("  H294（新增项8 审定金额）= %.6f" % H294)
P("  %-6s %-20s → %s" % ("", "旧缓存", "新缓存"))
for ref, o, n in [("H238", H238_OLD, H238_NEW), ("H213", H213_OLD, H213_NEW),
                  ("H4", H4_OLD, H4_NEW), ("H559", H559_OLD, H559_NEW)]:
    P("  %-6s %-20s → %.6f" % (ref, o, n))
P("  H558（施工单位报送复核基准，常量）= %.2f（不动）" % H558)
P("")

PATCHES = [
    ("H238", NEW_H238_F, H238_NEW),
    ("H213", None,       H213_NEW),
    ("H4",   None,       H4_NEW),
    ("H559", None,       H559_NEW),
]

# ---------- 选择落盘目标 ----------
def writable(p):
    try:
        with open(p, "r+b"):
            pass
        return True
    except Exception:
        return False

if writable(SRC):
    TARGET = SRC
    P(">> 源表可写，采用【原地修正】")
else:
    TARGET = CPY
    P(">> 源表被 WPS/Excel 占用，采用【副本修正】：%s" % os.path.basename(CPY))
    P("   （原文不动；关闭 WPS 后重跑本脚本可改为原地修正）")
    if not writable(CPY) and os.path.exists(CPY):
        P("!! 副本也被占用，中止"); open(OUT, "w", encoding="utf-8").write("\n".join(L)); print("BOTH LOCKED"); raise SystemExit(0)
P("")

# ---------- 取原 zip 内容（源表始终可读） ----------
zin = zipfile.ZipFile(SRC)
items = [(i, zin.read(i.filename)) for i in zin.infolist()]
zin.close()
xml = next(d for i, d in items if i.filename == SHEET_XML).decode("utf-8")

m = re.search(r'<c r="H238"[^>]*>.*?</c>', xml, re.S)
if NEW_H238_F in m.group(0):
    P("H238 公式已含 H294，无需修改。")
    open(OUT, "w", encoding="utf-8").write("\n".join(L)); print("ALREADY"); raise SystemExit(0)
if OLD_H238_F not in m.group(0):
    P("!! H238 公式与预期不符，中止：%s" % m.group(0)[:200])
    open(OUT, "w", encoding="utf-8").write("\n".join(L)); print("UNEXPECTED"); raise SystemExit(1)

P("=== 外科式改 XML ===")
for ref, newf, newv in PATCHES:
    mm = re.search(r'<c r="%s"[^>]*>.*?</c>' % ref, xml, re.S)
    if not mm:
        P("  !! 未找到 %s" % ref); continue
    old = mm.group(0)
    s_attr = re.search(r's="(\d+)"', old)
    oldv = re.search(r'<v>([^<]*)</v>', old)
    oldf = re.search(r'<f>([^<]*)</f>', old)
    f_use = newf if newf else oldf.group(1)
    nv = ("%.10f" % newv).rstrip("0").rstrip(".")
    xml = xml.replace(old, '<c r="%s" s="%s"><f>%s</f><v>%s</v></c>' % (
        ref, s_attr.group(1), f_use, nv), 1)
    P("  %-6s 公式 %-46s → %-46s" % (ref, oldf.group(1)[:44], f_use[:44]))
    P("          缓存 %-22s → %s" % (oldv.group(1) if oldv else "-", nv))

with zipfile.ZipFile(TARGET, "w", zipfile.ZIP_DEFLATED) as zout:
    for info, data in items:
        zout.writestr(info, xml.encode("utf-8") if info.filename == SHEET_XML else data)
P("  已写入：%s" % os.path.basename(TARGET))
P("  其余条目与缓存原样保留")
P("")

# ---------- 回读验证 ----------
ws = openpyxl.load_workbook(TARGET, data_only=True)["Sheet2"]
ls_g = ls_h = 0.0
for r in range(4, ws.max_row + 1):
    if isinstance(ws.cell(r, 6).value, (int, float)) and ws.cell(r, 2).value:
        ls_g += ws.cell(r, 7).value or 0.0
        ls_h += ws.cell(r, 8).value or 0.0
P("=== 修正后验证（对 TARGET 回读）===")
P("  H238 = %.6f   H213 = %.6f" % (ws.cell(238, 8).value, ws.cell(213, 8).value))
P("  H4   = %.6f   H559 = %.6f" % (ws.cell(4, 8).value, ws.cell(559, 8).value))
P("  G4   = %.2f（报送总，未动）" % ws.cell(4, 7).value)
P("  叶子求和 G = %.6f  vs G4 = %.6f  缺口 = %.6f" % (ls_g, ws.cell(4, 7).value, ls_g - (ws.cell(4, 7).value or 0)))
P("  叶子求和 H = %.6f  vs H4 = %.6f  缺口 = %.6f" % (ls_h, ws.cell(4, 8).value, ls_h - (ws.cell(4, 8).value or 0)))
P("  >>> 审定口径：%.6f（= 目标合并汇总 r103 校验值）" % ws.cell(4, 8).value)
open(OUT, "w", encoding="utf-8").write("\n".join(L))
print("done -> %s" % TARGET)
